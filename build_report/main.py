import os
import sys

# Ensure repository root is on sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from build_report.styles import (
    FONT_NAME, COLOR_PRIMARY, COLOR_TEXT, COLOR_MUTED,
    set_run_font, add_page_number_field
)
from build_report.preliminary import build_preliminary_pages
from build_report.chapters_1_3 import build_chapters_1_to_3
from build_report.chapters_4_5 import build_chapters_4_to_5
from build_report.chapters_6_7 import build_chapters_6_to_7
from build_report.chapters_8_10 import build_chapters_8_to_10
from build_report.chapters_11_14 import build_chapters_11_to_14
from build_report.backmatter import build_backmatter

def create_full_blackbook():
    print("=" * 60)
    print(">> STARTING EXPENSEX ACADEMIC BLACK BOOK GENERATION...")
    print("=" * 60)

    doc = docx.Document()

    # Configure Default Style
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = FONT_NAME
    font.size = Pt(12)
    font.color.rgb = COLOR_TEXT

    # =========================================================================
    # SECTION 1: PRELIMINARY PAGES (Roman Numerals)
    # =========================================================================
    sec1 = doc.sections[0]
    sec1.top_margin = Inches(1.0)
    sec1.bottom_margin = Inches(1.0)
    sec1.left_margin = Inches(1.25)   # Extra margin for binding / gutter
    sec1.right_margin = Inches(1.0)
    sec1.different_first_page_header_footer = True  # No header/footer on title page

    # Preliminary Footer for subsequent preliminary pages
    footer1 = sec1.footer
    p_f1 = footer1.paragraphs[0]
    p_f1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_f1.paragraph_format.space_before = Pt(4)
    p_f1.paragraph_format.space_after = Pt(0)
    r_f1 = p_f1.add_run("Preliminary — ")
    set_run_font(r_f1, name=FONT_NAME, size_pt=9.5, italic=True, color=COLOR_MUTED)
    add_page_number_field(p_f1.add_run())

    # Set Roman Numeral Format for Section 1
    sectPr1 = sec1._sectPr
    pgNumType1 = OxmlElement('w:pgNumType')
    pgNumType1.set(qn('w:fmt'), 'lowerRoman')
    sectPr1.append(pgNumType1)

    logo_path = os.path.join(os.path.dirname(__file__), "..", "client", "public", "logo1.png")
    logo_path = os.path.abspath(logo_path)

    print("[-] Building Preliminary Pages (Title, Cert, Decl, Ack, Abstract, TOC, LOF, LOT)...")
    build_preliminary_pages(doc, logo_path=logo_path)

    # =========================================================================
    # SECTION 2: MAIN CHAPTERS & BACKMATTER (Arabic Numerals starting at 1)
    # =========================================================================
    sec2 = doc.add_section()
    sec2.top_margin = Inches(1.0)
    sec2.bottom_margin = Inches(1.0)
    sec2.left_margin = Inches(1.25)
    sec2.right_margin = Inches(1.0)
    sec2.header.is_linked_to_previous = False
    sec2.footer.is_linked_to_previous = False

    # Section 2 Header
    header2 = sec2.header
    p_h2 = header2.paragraphs[0]
    p_h2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_h2.paragraph_format.space_after = Pt(4)
    r_h2 = p_h2.add_run("ExpenseX — Smart Expense Tracker | University Project Report")
    set_run_font(r_h2, name=FONT_NAME, size_pt=9, italic=True, color=COLOR_MUTED)

    # Section 2 Footer
    footer2 = sec2.footer
    p_f2 = footer2.paragraphs[0]
    p_f2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_f2.paragraph_format.space_before = Pt(4)
    p_f2.paragraph_format.space_after = Pt(0)
    
    # Left text: Project Report | B.Sc. Computer Science | [Academic Year]
    r_f2_left = p_f2.add_run("Project Report | B.Sc. Computer Science | [Academic Year]          ")
    set_run_font(r_f2_left, name=FONT_NAME, size_pt=9.5, color=COLOR_MUTED)
    
    # Tab to Right for Page Number
    r_tab = p_f2.add_run("\t\t\t\tPage ")
    set_run_font(r_tab, name=FONT_NAME, size_pt=9.5, color=COLOR_MUTED)
    add_page_number_field(p_f2.add_run())

    # Set Decimal Format and Restart at Page 1
    sectPr2 = sec2._sectPr
    pgNumType2 = OxmlElement('w:pgNumType')
    pgNumType2.set(qn('w:fmt'), 'decimal')
    pgNumType2.set(qn('w:start'), '1')
    sectPr2.append(pgNumType2)

    print("[-] Building Chapters 1, 2, and 3 (Preliminary Investigation, System Analysis, Existing vs Proposed)...")
    build_chapters_1_to_3(doc)

    print("[-] Building Chapters 4 and 5 (System Design, Diagrams, Database Architecture)...")
    build_chapters_4_to_5(doc)

    print("[-] Building Chapters 6 and 7 (System Implementation, Source Code Architecture, Testing Suite)...")
    build_chapters_6_to_7(doc)

    print("[-] Building Chapters 8, 9, and 10 (Security Measures, Gemini AI Integration, Screen Layouts & Results)...")
    build_chapters_8_to_10(doc)

    print("[-] Building Chapters 11, 12, 13, and 14 (Deployment, Project Management, Limitations & Future Scope, Conclusion)...")
    build_chapters_11_to_14(doc)

    print("[-] Building Backmatter (IEEE References, 35+ Glossary Terms, Appendices A through H)...")
    build_backmatter(doc)

    output_filename = "ExpenseX_Project_Report_BlackBook.docx"
    output_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", output_filename))
    
    print(f"[-] Saving complete Black Book document to: {output_path} ...")
    doc.save(output_path)

    # Calculate statistics
    total_paragraphs = len(doc.paragraphs)
    total_tables = len(doc.tables)
    total_words = sum(len(p.text.split()) for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                total_words += sum(len(p.text.split()) for p in cell.paragraphs)

    print("=" * 60)
    print("SUCCESS: EXPENSEX BLACK BOOK GENERATION COMPLETE!")
    print(f"Document File: {output_path}")
    print(f"Total Word Count: ~{total_words:,} words")
    print(f"Total Paragraphs: {total_paragraphs}")
    print(f"Total Formatted Tables: {total_tables}")
    print("=" * 60)

if __name__ == "__main__":
    create_full_blackbook()
